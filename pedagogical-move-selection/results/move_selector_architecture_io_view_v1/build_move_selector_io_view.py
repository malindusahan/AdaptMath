"""Build a read-only Excel view of the current LIVE Turn-LinTS architecture.

The script reads the authoritative LIVE lineage and passive turn-outcome data,
but never imports or calls the policy runtime. Source hashes are checked before
and after generation so a concurrent Tutor update cannot produce a mixed view.
"""

from __future__ import annotations

import hashlib
import json
import math
import os
from collections import Counter
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


SCRIPT_DIR = Path(__file__).resolve().parent
WORKSPACE_ROOT = Path(__file__).resolve().parents[3]
RUNTIME_ROOT = (
    WORKSPACE_ROOT
    / "adaptive-math-tutor"
    / "backend"
    / "runtime"
    / "turn_lints_live_skl_final_v1"
)
EVENT_PATH = RUNTIME_ROOT / "turn_events.jsonl"
STATE_PATH = RUNTIME_ROOT / "policy_state.json"
LINEAGE_PATH = RUNTIME_ROOT / "lineage_manifest.json"
PASSIVE_PATH = (
    WORKSPACE_ROOT
    / "pedagogical-move-selection"
    / "results"
    / "md_self_improvement_turn_data_v1"
    / "turn_outcomes.jsonl"
)
WORKBOOK_PATH = SCRIPT_DIR / "move_selector_architecture_io_view.xlsx"
SUMMARY_PATH = SCRIPT_DIR / "summary.json"

ARMS = ("generic", "probing", "focus", "telling")
FEATURES = (
    "selector_p_generic",
    "selector_p_probing",
    "selector_p_focus",
    "selector_p_telling",
    "mastery_before",
    "previous_mastery_delta",
    "previous_mastery_delta_missing",
    "previous_reasoning_probability",
    "previous_uncertainty_probability",
    "previous_clarification_probability",
    "previous_learner_signals_missing",
)
EXPECTED_SELECTOR_SHA = (
    "ad770256a71efcce7485705099511494e41806b2e29bc3829f9ba8aeaf731cd5"
)
EXPECTED_SELECTOR = "MD7-R2-TELL-C1 Epoch 2"
EXPECTED_SCHEMA = "adaptmath_turn_lints_event_v4"

COLORS = {
    "title": "17365D",
    "identity": "D9EAF7",
    "S": "BDD7EE",
    "K": "C6E0B4",
    "L": "F8CBAD",
    "decision": "D9E1F2",
    "output": "FFE699",
    "audit": "E2F0D9",
    "meta": "E7E6E6",
    "selected": "FFF2CC",
    "good": "C6E0B4",
    "bad": "F4CCCC",
    "white": "FFFFFF",
}
THIN = Side(style="thin", color="B7B7B7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def relative(path: Path) -> str:
    return path.resolve().relative_to(WORKSPACE_ROOT).as_posix()


def load_json(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError(f"{path} must contain one JSON object.")
    return value


def load_jsonl(path: Path) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for line_number, raw in enumerate(
        path.read_text(encoding="utf-8").splitlines(), start=1
    ):
        if not raw.strip():
            continue
        try:
            value = json.loads(raw)
        except json.JSONDecodeError as exc:
            raise ValueError(f"{path} line {line_number} is invalid JSON.") from exc
        if not isinstance(value, dict):
            raise ValueError(f"{path} line {line_number} is not an object.")
        records.append(value)
    return records


def nested(record: Mapping[str, Any] | None, *keys: str, default: Any = None) -> Any:
    current: Any = record
    for key in keys:
        if not isinstance(current, Mapping) or key not in current:
            return default
        current = current[key]
    return current


def finite_number(value: Any) -> bool:
    return (
        isinstance(value, (int, float))
        and not isinstance(value, bool)
        and math.isfinite(float(value))
    )


def json_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def md7_argmax(probabilities: Mapping[str, Any]) -> str:
    return max(ARMS, key=lambda arm: float(probabilities[arm]))


def normalized_status(event: Mapping[str, Any], passive: Mapping[str, Any] | None) -> str:
    passive_status = nested(passive, "next_learner_observation", "status")
    if isinstance(passive_status, str) and passive_status:
        return passive_status
    return {
        "observed_bkt_update": "observed_update",
        "no_bkt_update": "resolved_no_update",
        "censored_no_response": "censored_no_response",
        "processing_failed": "processing_failed",
    }.get(str(event.get("resolution_status")), str(event.get("resolution_status") or ""))


def recompute_reward(event: Mapping[str, Any]) -> tuple[float | None, str, float | None]:
    status = event.get("resolution_status")
    before = event.get("mastery_before")
    after = event.get("mastery_after")
    logged = event.get("configured_reward_value")
    if status != "observed_bkt_update" or not (
        finite_number(before) and finite_number(after)
    ):
        return None, "NOT_APPLICABLE", None
    before_f = float(before)
    after_f = float(after)
    delta = after_f - before_f
    if delta > 0.0:
        denominator = 1.0 - before_f
        reward = delta / denominator if denominator > 0.0 else 0.0
    elif delta < 0.0:
        reward = delta / before_f if before_f > 0.0 else 0.0
    else:
        reward = 0.0
    reward = max(-1.0, min(1.0, float(reward)))
    if not finite_number(logged):
        return reward, "DERIVED_FOR_VIEW_ONLY", None
    difference = float(logged) - reward
    return reward, "MATCH" if abs(difference) <= 1e-12 else "MISMATCH", difference


def update_eligibility(event: Mapping[str, Any]) -> tuple[bool, bool, str]:
    agency = bool(nested(event, "explicit_learner_agency", "triggered", default=False))
    selected = event.get("selected_arm") in ARMS
    policy_selected = bool(
        selected
        and event.get("treatment_assignment_source") == "turn_lints_posterior_sample"
        and not agency
    )
    if event.get("policy_mode") != "LIVE" or event.get("real_live_observation") is not True:
        return policy_selected, False, "not a real LIVE action"
    if agency or event.get("agency_forced_no_policy_update") is True:
        return policy_selected, False, "explicit learner-agency action; posterior update prohibited"
    if not policy_selected:
        return policy_selected, False, "action was not selected by the LIVE Turn-LinTS policy"
    status = event.get("resolution_status")
    if status == "censored_no_response":
        return policy_selected, False, "censored: no learner response"
    if status == "processing_failed":
        return policy_selected, False, "post-action processing failed"
    if status == "no_bkt_update":
        return policy_selected, False, "resolved learner response produced no valid BKT update"
    reward, check, _ = recompute_reward(event)
    if status != "observed_bkt_update" or reward is None:
        return policy_selected, False, "no valid linked BKT transition/reward"
    if check == "MISMATCH":
        return policy_selected, False, "logged reward does not match recomputation"
    return policy_selected, True, "eligible: policy-selected LIVE action with valid linked BKT reward"


def build_row(event: Mapping[str, Any], passive: Mapping[str, Any] | None) -> dict[str, Any]:
    context = nested(event, "context", default={})
    features = nested(context, "feature_values", default={})
    probabilities = event.get("md7_raw_probabilities") or {
        arm: nested(features, f"selector_p_{arm}") for arm in ARMS
    }
    policy_scores = event.get("policy_scores") or {}
    reward_scores = event.get("contextual_ts_reward_scores") or {}
    behavior_vector = event.get("full_behavior_probability_vector") or {}
    passive_decision = nested(passive, "adaptive_decision", default={})
    state_before = nested(passive, "state_before_action", default={})
    learning = nested(passive, "learning_state_update", default={})
    learner = nested(passive, "next_learner_observation", default={})
    resolver = nested(learner, "resolver", default={})
    tutor_quality = nested(passive, "tutor_quality", default={})
    agency = nested(event, "explicit_learner_agency", default={})
    recomputed, reward_check, reward_difference = recompute_reward(event)
    policy_selected, eligible, eligibility_reason = update_eligibility(event)
    status = normalized_status(event, passive)
    vector = list(nested(context, "vector", default=[]))
    contributors = nested(resolver, "contributors", default=[])
    timestamp = nested(passive, "provenance", "timestamp") or event.get("update_timestamp")
    row: dict[str, Any] = {
        "timestamp": timestamp,
        "action_event_id": event.get("action_event_id"),
        "attempt_id": event.get("attempt_id"),
        "turn_index": event.get("action_turn_index"),
        "skill": nested(state_before, "target_skill"),
        "problem": nested(state_before, "problem"),
        "latest_student_text_before_action": nested(state_before, "latest_student_text"),
        **{f"p_{arm}": nested(probabilities, arm) for arm in ARMS},
        "MD7_argmax": md7_argmax(probabilities),
        "mastery_before": nested(features, "mastery_before"),
        "previous_mastery_delta": nested(features, "previous_mastery_delta"),
        "previous_mastery_delta_missing": nested(features, "previous_mastery_delta_missing"),
        "previous_reasoning_probability": nested(features, "previous_reasoning_probability"),
        "previous_uncertainty_probability": nested(features, "previous_uncertainty_probability"),
        "previous_clarification_probability": nested(features, "previous_clarification_probability"),
        "previous_learner_signals_missing": nested(features, "previous_learner_signals_missing"),
        **{f"score_{arm}": nested(policy_scores, arm) for arm in ARMS},
        "selected_arm": event.get("selected_arm"),
        "final_move": event.get("actual_move") or event.get("actual_treatment_move"),
        "decision_source": event.get("decision_source"),
        "treatment_assignment_source": event.get("treatment_assignment_source"),
        "explicit_learner_agency_triggered": nested(agency, "triggered", default=False),
        "tutor_response": event.get("tutor_response") or nested(passive, "tutor_response"),
        "learner_response": nested(learner, "student_text"),
        "resolution_status": status,
        "mastery_after": event.get("mastery_after"),
        "delta_mastery": event.get("raw_delta"),
        "headroom_normalized_reward": event.get("headroom_normalized_delta"),
        "posterior_should_update": eligible,
        "posterior_update_weight": event.get("posterior_update_weight"),
        "policy_mode": event.get("policy_mode"),
        "policy_lineage": event.get("policy_lineage"),
        "behavior_policy": event.get("behavior_policy"),
        "context_schema": nested(context, "schema_id"),
        "context_dimension": nested(context, "context_dimension"),
        "reward_mode": event.get("reward_mode"),
        "context_vector_json": json_cell(vector),
        "available_arms": json_cell(nested(passive_decision, "available_arms", default=list(ARMS))),
        "base_move": event.get("actual_selector_move") or nested(passive_decision, "base_move"),
        "randomized_assignment": event.get("randomized_assignment"),
        "behavior_propensity": event.get("behavior_propensity"),
        **{f"behavior_p_{arm}": nested(behavior_vector, arm) for arm in ARMS},
        "behavior_probability_vector": json_cell(behavior_vector),
        "learner_agency_category": nested(agency, "request_category"),
        "learner_agency_source": (
            event.get("treatment_assignment_source") if nested(agency, "triggered", default=False) else None
        ),
        **{
            f"mrb1_{name}": nested(event, "current_response_mrb1", name)
            for name in (
                "Mistake_Identification",
                "Mistake_Location",
                "Providing_Guidance",
                "Actionability",
            )
        },
        "mrb1_model_version": nested(tutor_quality, "model_version"),
        "mrb1_model_sha256": nested(tutor_quality, "model_sha256"),
        "bkt_should_update": nested(learning, "should_update"),
        "bkt_observation_source": nested(learning, "observation_source"),
        "bkt_outcome": nested(learning, "outcome"),
        "bkt_update_confidence": nested(learning, "update_confidence"),
        "resolver_primary_signal": nested(resolver, "primary_signal"),
        "resolver_evidence_weight": nested(resolver, "evidence_weight"),
        "resolver_contributors": json_cell(contributors),
        "logged_reward": event.get("configured_reward_value"),
        "recomputed_reward": recomputed,
        "reward_difference": reward_difference,
        "reward_check": reward_check,
        "posterior_updated": event.get("posterior_updated"),
        "posterior_update_origin": event.get("posterior_update_origin"),
        "event_resolution_status": event.get("resolution_status"),
        "was_policy_selected": policy_selected,
        "eligible_for_LinTS_update": eligible,
        "eligibility_reason": eligibility_reason,
        "actual_selector_move": event.get("actual_selector_move"),
        "actual_treatment_move": event.get("actual_treatment_move"),
        "configured_reward_mode": event.get("configured_reward_mode"),
        "outcome_observed": event.get("outcome_observed"),
        "real_live_observation": event.get("real_live_observation"),
        "posterior_update_occurred": event.get("posterior_update_occurred"),
    }
    for index, feature in enumerate(FEATURES, start=1):
        row[f"ctx_{index:02d}_{feature}"] = vector[index - 1] if len(vector) >= index else None
    for arm in ARMS:
        row[f"reward_score_{arm}"] = nested(reward_scores, arm)
    return row


DECISION_COLUMNS: list[tuple[str, str, str]] = [
    ("timestamp", "Timestamp (pre-action record)", "identity"),
    ("action_event_id", "Action event ID", "identity"),
    ("attempt_id", "Attempt ID", "identity"),
    ("turn_index", "Turn index", "identity"),
    ("skill", "Canonical BKT skill", "identity"),
    ("problem", "Problem", "identity"),
    ("latest_student_text_before_action", "Latest student text before action", "identity"),
    *[(f"p_{arm}", f"p_{arm}", "S") for arm in ARMS],
    ("MD7_argmax", "MD7 argmax", "S"),
    ("mastery_before", "Mastery before", "K"),
    ("previous_mastery_delta", "Previous mastery delta", "K"),
    ("previous_mastery_delta_missing", "Previous mastery delta missing", "K"),
    ("previous_reasoning_probability", "Previous reasoning probability", "L"),
    ("previous_uncertainty_probability", "Previous uncertainty probability", "L"),
    ("previous_clarification_probability", "Previous clarification probability", "L"),
    ("previous_learner_signals_missing", "Previous learner signals missing", "L"),
    *[(f"score_{arm}", f"Score {arm}", "decision") for arm in ARMS],
    ("selected_arm", "Selected arm", "decision"),
    ("final_move", "Final move", "decision"),
    ("decision_source", "Decision source", "decision"),
    ("treatment_assignment_source", "Treatment assignment source", "decision"),
    ("explicit_learner_agency_triggered", "Explicit learner agency triggered", "decision"),
    ("tutor_response", "Tutor response", "output"),
    ("learner_response", "Learner response", "output"),
    ("resolution_status", "Resolution status", "output"),
    ("mastery_after", "Mastery after", "output"),
    ("delta_mastery", "Delta mastery", "output"),
    ("headroom_normalized_reward", "HEADROOM_NORMALIZED reward", "output"),
    ("posterior_should_update", "Posterior should update", "output"),
    ("posterior_update_weight", "Posterior update weight", "output"),
]

FORENSIC_EXTRA_COLUMNS: list[tuple[str, str, str]] = [
    ("policy_mode", "Policy mode", "meta"),
    ("policy_lineage", "Policy lineage", "meta"),
    ("behavior_policy", "Behavior policy", "meta"),
    ("context_schema", "Context schema", "meta"),
    ("context_dimension", "Context dimension", "meta"),
    ("reward_mode", "Reward mode", "meta"),
    ("context_vector_json", "Full context vector (JSON)", "meta"),
    *[(f"ctx_{i:02d}_{name}", f"Context {i:02d}: {name}", "meta") for i, name in enumerate(FEATURES, 1)],
    *[(f"reward_score_{arm}", f"Reward score {arm}", "decision") for arm in ARMS],
    ("available_arms", "Available arms", "decision"),
    ("base_move", "Base move / MD7 recommendation", "decision"),
    ("randomized_assignment", "Randomized assignment", "decision"),
    ("behavior_propensity", "Behavior propensity", "decision"),
    *[(f"behavior_p_{arm}", f"Behavior p {arm}", "decision") for arm in ARMS],
    ("behavior_probability_vector", "Full behavior probability vector", "decision"),
    ("learner_agency_category", "Learner-agency category", "decision"),
    ("learner_agency_source", "Learner-agency source", "decision"),
    ("mrb1_Mistake_Identification", "MRB1 Mistake Identification", "audit"),
    ("mrb1_Mistake_Location", "MRB1 Mistake Location", "audit"),
    ("mrb1_Providing_Guidance", "MRB1 Providing Guidance", "audit"),
    ("mrb1_Actionability", "MRB1 Actionability", "audit"),
    ("mrb1_model_version", "MRB1 model version", "audit"),
    ("mrb1_model_sha256", "MRB1 model SHA256", "audit"),
    ("bkt_should_update", "BKT should update", "audit"),
    ("bkt_observation_source", "BKT observation source", "audit"),
    ("bkt_outcome", "BKT outcome", "audit"),
    ("bkt_update_confidence", "BKT update confidence", "audit"),
    ("resolver_primary_signal", "Resolver primary signal", "audit"),
    ("resolver_evidence_weight", "Resolver evidence weight", "audit"),
    ("resolver_contributors", "Resolver contributors", "audit"),
    ("event_resolution_status", "Event resolution status", "audit"),
    ("logged_reward", "Logged reward", "audit"),
    ("recomputed_reward", "Recomputed reward", "audit"),
    ("reward_difference", "Reward difference", "audit"),
    ("reward_check", "Reward check", "audit"),
    ("posterior_updated", "Posterior updated", "audit"),
    ("posterior_update_occurred", "Posterior update occurred", "audit"),
    ("posterior_update_origin", "Posterior update origin", "audit"),
    ("was_policy_selected", "Was policy selected", "audit"),
    ("eligible_for_LinTS_update", "Eligible for LinTS update", "audit"),
    ("eligibility_reason", "Eligibility reason", "audit"),
]


def style_title(ws: Any, title: str, note: str, max_col: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    cell = ws.cell(1, 1, title)
    cell.font = Font(size=16, bold=True, color=COLORS["white"])
    cell.fill = PatternFill("solid", fgColor=COLORS["title"])
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_col)
    note_cell = ws.cell(2, 1, note)
    note_cell.alignment = Alignment(wrap_text=True, vertical="top")
    note_cell.fill = PatternFill("solid", fgColor="F2F2F2")
    ws.row_dimensions[2].height = 34
    ws.sheet_view.showGridLines = False


def write_grouped_table(
    ws: Any,
    title: str,
    note: str,
    columns: Sequence[tuple[str, str, str]],
    rows: Sequence[Mapping[str, Any]],
) -> None:
    style_title(ws, title, note, len(columns))
    start_col = 1
    while start_col <= len(columns):
        group = columns[start_col - 1][2]
        end_col = start_col
        while end_col < len(columns) and columns[end_col][2] == group:
            end_col += 1
        ws.merge_cells(start_row=3, start_column=start_col, end_row=3, end_column=end_col)
        group_cell = ws.cell(3, start_col, group.upper())
        group_cell.fill = PatternFill("solid", fgColor=COLORS.get(group, COLORS["meta"]))
        group_cell.font = Font(bold=True)
        group_cell.alignment = Alignment(horizontal="center")
        start_col = end_col + 1
    for col_index, (_, heading, group) in enumerate(columns, start=1):
        cell = ws.cell(4, col_index, heading)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=COLORS.get(group, COLORS["meta"]))
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER
    ws.row_dimensions[4].height = 44
    numeric_tokens = (
        "probability", "mastery", "delta", "reward", "score", "confidence",
        "propensity", "weight", "p_generic", "p_probing", "p_focus", "p_telling",
        "ctx_", "behavior_p_",
    )
    for row_index, row in enumerate(rows, start=5):
        selected_arm = row.get("selected_arm")
        for col_index, (key, _, group) in enumerate(columns, start=1):
            value = row.get(key)
            cell = ws.cell(row_index, col_index, value)
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if finite_number(value) and any(token in key.lower() for token in numeric_tokens):
                cell.number_format = "0.0000"
            if key == f"score_{selected_arm}" or key == "selected_arm":
                cell.fill = PatternFill("solid", fgColor=COLORS["selected"])
                cell.font = Font(bold=True)
            if key in {"reward_check", "eligible_for_LinTS_update", "posterior_should_update"}:
                if value in {"MATCH", True}:
                    cell.fill = PatternFill("solid", fgColor=COLORS["good"])
                elif value in {"MISMATCH", False}:
                    cell.fill = PatternFill("solid", fgColor=COLORS["bad"])
        ws.row_dimensions[row_index].height = 42
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(columns))}{4 + len(rows)}"
    text_wide = {
        "problem": 42,
        "latest_student_text_before_action": 34,
        "tutor_response": 58,
        "learner_response": 42,
        "action_event_id": 49,
        "attempt_id": 40,
        "mrb1_model_sha256": 35,
        "context_vector_json": 45,
        "behavior_probability_vector": 38,
        "eligibility_reason": 45,
    }
    for col_index, (key, heading, _) in enumerate(columns, start=1):
        width = text_wide.get(key, min(28, max(12, len(heading) + 2)))
        ws.column_dimensions[get_column_letter(col_index)].width = width


def add_section(ws: Any, row: int, title: str, width: int = 4) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
    cell = ws.cell(row, 1, title)
    cell.font = Font(bold=True, color=COLORS["white"])
    cell.fill = PatternFill("solid", fgColor=COLORS["title"])
    return row + 1


def add_key_value(ws: Any, row: int, key: str, value: Any, note: str = "") -> int:
    for col, item in enumerate((key, value, note), start=1):
        cell = ws.cell(row, col, item)
        cell.border = BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if col == 1:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
    return row + 1


def build_summary_sheet(
    wb: Workbook,
    state: Mapping[str, Any],
    lineage: Mapping[str, Any],
    counts: Mapping[str, Any],
    source_hashes: Mapping[str, str],
    accounting_status: str,
) -> None:
    ws = wb.active
    ws.title = "Summary"
    style_title(
        ws,
        "Current LIVE Move-Selector Architecture",
        "Read-only snapshot of what the active policy receives before each Tutor action and what follows after it.",
        4,
    )
    row = 4
    row = add_section(ws, row, "Architecture")
    architecture_items = [
        ("Architecture", "AdaptMath Turn-LinTS v1", ""),
        ("Runtime mode", lineage["mode"], ""),
        ("Algorithm", "direct disjoint Turn-LinTS", lineage["algorithm"]),
        ("Arms", ", ".join(ARMS), ""),
        ("Policy context", lineage["context_preset"], "S + K + L"),
        ("Context dimension", state["context_dimension"], "Exactly 11 policy inputs"),
        ("Reward", "HEADROOM_NORMALIZED", state["reward_mode"]),
        ("Selector", state["selector_version"], "Frozen base selector"),
        ("Selector SHA", state["selector_sha256"], ""),
        ("Anchor", state["anchor_mode"], f"gamma={state['anchor_gamma']}"),
        ("Tau", "none", "No tau gate"),
        ("S", "Frozen MD7 probabilities", "4 features"),
        ("K", "Knowledge state", "3 features"),
        ("L", "Previous learner-response state", "4 features"),
        ("Diagnostic only", "H + Q / MRB1 where logged", "Not used by final S+K+L policy"),
    ]
    for item in architecture_items:
        row = add_key_value(ws, row, *item)
    row += 1
    row = add_section(ws, row, "Current LIVE Counts")
    count_items = [
        ("Total LIVE actions", counts["total_live_actions"], "One row per actual Tutor action"),
        ("observed_update", counts["observed_update"], "Valid linked BKT transition"),
        ("resolved_no_update", counts["resolved_no_update"], "Resolved but no BKT update/reward"),
        ("censored_no_response", counts["censored_no_response"], "No learner response"),
        ("processing_failed", counts["processing_failed"], "Post-action processing failure"),
        *[(f"{arm} selections", counts["move_counts"][arm], "") for arm in ARMS],
        ("Current policy_state update_count", state["update_count"], "Synthetic plus real LIVE update identities"),
        ("Synthetic initialization count", lineage["synthetic_observation_count"], "From lineage manifest"),
        ("Synthetic effective weight", lineage["synthetic_effective_sample_size"], "From lineage manifest"),
        ("Valid real LIVE updates from log", counts["real_live_updates"], "posterior_updated=true events"),
        ("State / log accounting", accounting_status, "Compared using manifest provenance and per-arm counters"),
    ]
    for item in count_items:
        row = add_key_value(ws, row, *item)
    row += 1
    row = add_section(ws, row, "Authoritative Sources")
    for label, path, hash_key in (
        ("LIVE event log", EVENT_PATH, "event_log"),
        ("Policy state", STATE_PATH, "policy_state"),
        ("LIVE lineage manifest", LINEAGE_PATH, "lineage_manifest"),
        ("Passive text/BKT enrichment", PASSIVE_PATH, "passive_turns"),
    ):
        row = add_key_value(ws, row, label, relative(path), f"SHA256 {source_hashes[hash_key]}")
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 72
    ws.column_dimensions["C"].width = 65
    ws.column_dimensions["D"].width = 3
    ws.freeze_panes = "A4"
    ws.sheet_view.showGridLines = False


def build_field_guide(wb: Workbook) -> None:
    ws = wb.create_sheet("Field_Guide")
    headers = (
        "Field", "Block", "Stage", "Available before action?",
        "Used by final Turn-LinTS?", "Description",
    )
    style_title(
        ws,
        "Field Guide",
        "S = selector information\nK = learner knowledge state\nL = preceding learner-response state",
        len(headers),
    )
    rows = [
        ("selector_p_generic", "S", "Pre-action", "Yes", "Yes", "Frozen MD7 probability for generic."),
        ("selector_p_probing", "S", "Pre-action", "Yes", "Yes", "Frozen MD7 probability for probing."),
        ("selector_p_focus", "S", "Pre-action", "Yes", "Yes", "Frozen MD7 probability for focus."),
        ("selector_p_telling", "S", "Pre-action", "Yes", "Yes", "Frozen MD7 probability for telling."),
        ("mastery_before", "K", "Pre-action", "Yes", "Yes", "BKT mastery available when the move is selected."),
        ("previous_mastery_delta", "K", "Pre-action", "Yes", "Yes", "Mastery change caused by the preceding linked learner response; zero when missing."),
        ("previous_mastery_delta_missing", "K", "Pre-action", "Yes", "Yes", "1 when the preceding mastery delta is unavailable, otherwise 0."),
        ("previous_reasoning_probability", "L", "Pre-action", "Yes", "Yes", "Reasoning signal from the preceding learner response; zero when missing."),
        ("previous_uncertainty_probability", "L", "Pre-action", "Yes", "Yes", "Uncertainty signal from the preceding learner response; zero when missing."),
        ("previous_clarification_probability", "L", "Pre-action", "Yes", "Yes", "Clarification signal from the preceding learner response; zero when missing."),
        ("previous_learner_signals_missing", "L", "Pre-action", "Yes", "Yes", "1 when preceding learner signals are unavailable, otherwise 0."),
        ("score_generic", "Decision", "Action selection", "Computed now", "Yes", "Sampled Turn-LinTS score for generic."),
        ("score_probing", "Decision", "Action selection", "Computed now", "Yes", "Sampled Turn-LinTS score for probing."),
        ("score_focus", "Decision", "Action selection", "Computed now", "Yes", "Sampled Turn-LinTS score for focus."),
        ("score_telling", "Decision", "Action selection", "Computed now", "Yes", "Sampled Turn-LinTS score for telling."),
        ("selected_arm", "Decision", "Action selection", "Computed now", "Yes", "Argmax of the four sampled policy scores."),
        ("tutor_response", "Output", "Post-action", "No", "No", "Tutor text realized under the selected move contract."),
        ("learner_response", "Feedback", "Post-action", "No", "No", "Learner text immediately following the Tutor action."),
        ("mastery_after", "BKT", "Post-action", "No", "No", "BKT mastery after valid linked evidence."),
        ("delta_mastery", "BKT", "Post-action", "No", "No", "mastery_after minus mastery_before."),
        ("HEADROOM reward", "Reward", "Post-action", "No", "Used only to update", "Symmetric headroom-normalized mastery delta for the selected arm."),
        ("current MRB1", "Diagnostic Q", "Post-action", "No", "No", "Post-action Tutor-quality auxiliary output. It is not part of S+K+L and is not scalar reward."),
    ]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(4, col, header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=COLORS["meta"])
        cell.border = BORDER
        cell.alignment = Alignment(wrap_text=True)
    for row_index, values in enumerate(rows, 5):
        for col, value in enumerate(values, 1):
            cell = ws.cell(row_index, col, value)
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if col == 2 and value in {"S", "K", "L"}:
                cell.fill = PatternFill("solid", fgColor=COLORS[str(value)])
    widths = (38, 18, 20, 24, 26, 78)
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:F{4 + len(rows)}"
    ws.sheet_view.showGridLines = False


def build_architecture_flow(wb: Workbook) -> None:
    ws = wb.create_sheet("Architecture_Flow")
    headers = ("Phase", "Step", "Component", "Receives", "Produces / role")
    style_title(
        ws,
        "Architecture Flow",
        "The policy boundary is explicit: only S+K+L values available before selection enter the 11-D Turn-LinTS context.",
        len(headers),
    )
    flow = [
        ("PRE-ACTION", 1, "Learner/history state", "Prior conversation and persisted mastery", "State available before this action"),
        ("PRE-ACTION", 2, "Frozen MD7", "Problem and prior dialogue", "Four frozen selector probabilities"),
        ("PRE-ACTION", 3, "S", "p_generic, p_probing, p_focus, p_telling", "4 selector features"),
        ("PRE-ACTION", 4, "K", "mastery_before, previous mastery delta + missing flag", "3 knowledge features"),
        ("PRE-ACTION", 5, "L", "reasoning, uncertainty, clarification + missing flag", "4 preceding learner-state features"),
        ("PRE-ACTION", 6, "Concatenate S+K+L", "4 + 3 + 4 features", "Ordered 11-D policy vector"),
        ("PRE-ACTION", 7, "Direct disjoint Turn-LinTS", "11-D vector and four arm posteriors", "One sampled reward score per arm"),
        ("PRE-ACTION", 8, "Argmax sampled score", "generic/probing/focus/telling scores", "Selected pedagogical move"),
        ("POST-ACTION", 9, "Tutor realization contract", "Selected move and question context", "Tutor response constrained to the move"),
        ("POST-ACTION", 10, "Tutor response", "Realized response", "Learner-facing teaching turn"),
        ("POST-ACTION", 11, "Learner response", "Student text", "Evaluator/resolver evidence"),
        ("POST-ACTION", 12, "BKT", "Valid linked resolved evidence", "mastery_after and delta"),
        ("POST-ACTION", 13, "HEADROOM reward", "mastery_before and mastery_after", "Signed normalized reward in [-1, 1]"),
        ("POST-ACTION", 14, "Immediate selected-arm update", "Eligible policy action + valid reward", "Unit-weight A/b posterior update; otherwise no update"),
    ]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(4, col, header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=COLORS["meta"])
        cell.border = BORDER
    for row_index, values in enumerate(flow, 5):
        fill = COLORS["S"] if values[0] == "PRE-ACTION" else COLORS["output"]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row_index, col, value)
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if col == 1:
                cell.fill = PatternFill("solid", fgColor=fill)
                cell.font = Font(bold=True)
    for col, width in enumerate((18, 10, 34, 62, 62), 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False


def build_posterior_sheet(
    wb: Workbook,
    state: Mapping[str, Any],
    lineage: Mapping[str, Any],
    source_hash: str,
    log_update_counts: Mapping[str, int],
) -> None:
    ws = wb.create_sheet("Posterior_State")
    style_title(ws, "Current Posterior State", "Read-only inspection of policy_state.json; no posterior operation is executed.", 7)
    row = 4
    row = add_section(ws, row, "State Metadata", 7)
    metadata = [
        ("State path", relative(STATE_PATH), ""),
        ("State SHA256", source_hash, ""),
        ("Schema/version", state["schema_version"], state["algorithm_version"]),
        ("Context preset", "+".join(state["enabled_context_blocks"]), state["context_schema_id"]),
        ("Context dimension", state["context_dimension"], ""),
        ("Reward mode", state["reward_mode"], ""),
        ("Selector SHA", state["selector_sha256"], state["selector_version"]),
        ("Exploration", state["posterior_hyperparameters"]["exploration_scale"], ""),
        ("Lambda", state["posterior_hyperparameters"]["ridge_lambda"], ""),
        ("Update count", state["update_count"], ""),
        ("Synthetic initialization", lineage["synthetic_observation_count"], f"effective weight {lineage['synthetic_effective_sample_size']}"),
    ]
    for item in metadata:
        row = add_key_value(ws, row, *item)
    row += 1
    row = add_section(ws, row, "Per-Arm Posterior", 7)
    headers = (
        "Arm", "State updates", "Synthetic updates", "Derived real updates",
        "Logged real updates", "A shape", "A condition number / b length",
    )
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row, col, header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=COLORS["decision"])
        cell.border = BORDER
        cell.alignment = Alignment(wrap_text=True)
    row += 1
    theta_by_arm: dict[str, np.ndarray] = {}
    for arm in ARMS:
        matrix = np.asarray(state["A"][arm], dtype=float)
        vector = np.asarray(state["b"][arm], dtype=float)
        theta_by_arm[arm] = np.linalg.solve(matrix, vector)
        synthetic = int(lineage["synthetic_arm_update_counts"][arm])
        total = int(state["arm_update_counts"][arm])
        values = (
            arm,
            total,
            synthetic,
            total - synthetic,
            log_update_counts[arm],
            f"{matrix.shape[0]} x {matrix.shape[1]}",
            f"condition={np.linalg.cond(matrix):.6g}; b length={len(vector)}",
        )
        for col, value in enumerate(values, 1):
            cell = ws.cell(row, col, value)
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True)
        row += 1
    row += 1
    row = add_section(ws, row, "Posterior Mean Coefficients: theta_hat = A^-1 b", 7)
    coefficient_headers = ("Index", "Feature", *ARMS)
    for col, header in enumerate(coefficient_headers, 1):
        cell = ws.cell(row, col, header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=COLORS["decision"])
        cell.border = BORDER
    row += 1
    for index, feature in enumerate(state["ordered_feature_names"], 1):
        values = (index, feature, *(float(theta_by_arm[arm][index - 1]) for arm in ARMS))
        for col, value in enumerate(values, 1):
            cell = ws.cell(row, col, value)
            cell.border = BORDER
            if col >= 3:
                cell.number_format = "0.000000"
        row += 1
    for col, width in enumerate((10, 46, 20, 20, 20, 20, 45), 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A4"
    ws.sheet_view.showGridLines = False


def build_update_audit(
    wb: Workbook,
    rows: Sequence[Mapping[str, Any]],
    state: Mapping[str, Any],
    lineage: Mapping[str, Any],
    log_update_counts: Mapping[str, int],
) -> None:
    ws = wb.create_sheet("LIVE_Update_Audit")
    style_title(
        ws,
        "LIVE Update Audit",
        "Eligibility follows current runtime semantics: a policy-selected LIVE action needs a valid linked BKT reward and must not be learner-agency forced.",
        12,
    )
    ws.cell(4, 1, "Per-arm state/log accounting").font = Font(bold=True, color=COLORS["white"])
    ws.cell(4, 1).fill = PatternFill("solid", fgColor=COLORS["title"])
    headers = (
        "Arm", "Synthetic", "State total", "State-derived real", "Log real", "Status"
    )
    for col, header in enumerate(headers, 1):
        cell = ws.cell(5, col, header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=COLORS["decision"])
        cell.border = BORDER
    for index, arm in enumerate(ARMS, 6):
        synthetic = int(lineage["synthetic_arm_update_counts"][arm])
        total = int(state["arm_update_counts"][arm])
        derived = total - synthetic
        values = (arm, synthetic, total, derived, log_update_counts[arm], "MATCH" if derived == log_update_counts[arm] else "MISMATCH")
        for col, value in enumerate(values, 1):
            cell = ws.cell(index, col, value)
            cell.border = BORDER
            if col == 6:
                cell.fill = PatternFill("solid", fgColor=COLORS["good"] if value == "MATCH" else COLORS["bad"])
    action_headers = (
        "Action event ID", "Selected arm", "Resolution status", "Should update",
        "Computed headroom reward", "Was policy selected", "Learner agency triggered",
        "Eligible for LinTS update", "Reason", "Posterior updated",
        "Posterior update weight", "Accounting check",
    )
    header_row = 12
    for col, header in enumerate(action_headers, 1):
        cell = ws.cell(header_row, col, header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=COLORS["audit"])
        cell.border = BORDER
        cell.alignment = Alignment(wrap_text=True)
    for row_index, item in enumerate(rows, header_row + 1):
        expected = bool(item["eligible_for_LinTS_update"])
        actual = bool(item["posterior_updated"])
        values = (
            item["action_event_id"], item["selected_arm"], item["resolution_status"],
            item["bkt_should_update"], item["recomputed_reward"],
            item["was_policy_selected"], item["explicit_learner_agency_triggered"],
            expected, item["eligibility_reason"], actual,
            item["posterior_update_weight"], "MATCH" if expected == actual else "MISMATCH",
        )
        for col, value in enumerate(values, 1):
            cell = ws.cell(row_index, col, value)
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if col == 5 and finite_number(value):
                cell.number_format = "0.0000"
            if col == 12:
                cell.fill = PatternFill("solid", fgColor=COLORS["good"] if value == "MATCH" else COLORS["bad"])
    widths = (50, 18, 24, 18, 27, 22, 25, 25, 64, 20, 24, 20)
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:L{header_row + len(rows)}"
    ws.sheet_view.showGridLines = False


def build_historical_sheet(wb: Workbook, events: Sequence[Mapping[str, Any]]) -> None:
    if not events:
        return
    ws = wb.create_sheet("Historical_NonLIVE")
    columns = (
        "policy_mode", "action_event_id", "attempt_id", "action_turn_index",
        "selected_arm", "actual_move", "resolution_status", "update_timestamp",
    )
    style_title(ws, "Historical Non-LIVE Events", "Separated from the primary LIVE sheets.", len(columns))
    for col, key in enumerate(columns, 1):
        cell = ws.cell(4, col, key)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=COLORS["meta"])
        cell.border = BORDER
    for row_index, event in enumerate(events, 5):
        for col, key in enumerate(columns, 1):
            cell = ws.cell(row_index, col, event.get(key))
            cell.border = BORDER
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:H{4 + len(events)}"


def validate_inputs(
    state: Mapping[str, Any],
    lineage: Mapping[str, Any],
    architecture: Mapping[str, Any],
    events: Sequence[Mapping[str, Any]],
    rows: Sequence[Mapping[str, Any]],
) -> dict[str, Any]:
    errors: list[str] = []
    ids = [str(event.get("action_event_id")) for event in events]
    if len(ids) != len(set(ids)):
        errors.append("duplicate action_event_id in primary LIVE data")
    if tuple(state.get("ordered_feature_names", ())) != FEATURES:
        errors.append("policy state feature order is not the frozen 11-D order")
    if tuple(lineage.get("ordered_context_features", ())) != FEATURES:
        errors.append("lineage feature order is not the frozen 11-D order")
    if tuple(architecture.get("ordered_context_features", ())) != FEATURES:
        errors.append("architecture feature order is not the frozen 11-D order")
    if state.get("context_dimension") != 11 or lineage.get("context_dimension") != 11:
        errors.append("context dimension is not 11")
    if "mastery_after" in FEATURES:
        errors.append("mastery_after leaked into pre-action features")
    if any("mrb1" in name.lower() for name in FEATURES):
        errors.append("current MRB1 leaked into final policy context")
    for event, row in zip(events, rows, strict=True):
        if event.get("schema_version") != EXPECTED_SCHEMA:
            errors.append(f"unexpected event schema for {event.get('action_event_id')}")
        context = event.get("context") or {}
        if tuple(context.get("feature_names", ())) != FEATURES or len(context.get("vector", ())) != 11:
            errors.append(f"invalid 11-D context for {event.get('action_event_id')}")
        probabilities = [row[f"p_{arm}"] for arm in ARMS]
        if not all(finite_number(value) for value in probabilities):
            errors.append(f"non-finite MD7 probability for {event.get('action_event_id')}")
        elif abs(sum(float(value) for value in probabilities) - 1.0) > 1e-5:
            errors.append(f"MD7 probabilities do not sum to 1 for {event.get('action_event_id')}")
        if not all(finite_number(row[f"score_{arm}"]) for arm in ARMS):
            errors.append(f"missing Turn-LinTS score for {event.get('action_event_id')}")
        if row["selected_arm"] not in ARMS:
            errors.append(f"invalid selected arm for {event.get('action_event_id')}")
        if row["reward_check"] == "MISMATCH":
            errors.append(f"reward mismatch for {event.get('action_event_id')}")
        if bool(row["eligible_for_LinTS_update"]) != bool(row["posterior_updated"]):
            errors.append(f"eligibility/posterior mismatch for {event.get('action_event_id')}")
    if state.get("selector_sha256") != EXPECTED_SELECTOR_SHA:
        errors.append("selector SHA mismatch")
    if state.get("selector_version") != EXPECTED_SELECTOR:
        errors.append("selector version mismatch")
    if errors:
        raise ValueError("Workbook validation failed:\n- " + "\n- ".join(errors))
    return {
        "duplicate_action_event_ids": 0,
        "policy_context_feature_count": 11,
        "feature_order_matches_frozen_architecture": True,
        "probabilities_finite_and_sum_to_one": True,
        "four_scores_present": True,
        "selected_arms_valid": True,
        "mastery_after_excluded_from_pre_action_context": True,
        "current_mrb1_excluded_from_policy_context": True,
        "reward_recomputation_matches": True,
        "eligibility_matches_posterior_updates": True,
    }


def main() -> None:
    required = (EVENT_PATH, STATE_PATH, LINEAGE_PATH, PASSIVE_PATH)
    for path in required:
        if not path.is_file():
            raise FileNotFoundError(path)
    source_paths = {
        "event_log": EVENT_PATH,
        "policy_state": STATE_PATH,
        "lineage_manifest": LINEAGE_PATH,
        "passive_turns": PASSIVE_PATH,
    }
    source_hashes_before = {key: sha256(path) for key, path in source_paths.items()}
    state = load_json(STATE_PATH)
    lineage = load_json(LINEAGE_PATH)
    architecture_path = Path(lineage["architecture_manifest_path"])
    architecture = load_json(architecture_path)
    architecture_hash = sha256(architecture_path)
    if architecture_hash != str(lineage["architecture_manifest_sha256"]).lower():
        raise ValueError("Frozen architecture manifest hash does not match LIVE lineage.")
    source_hashes_before["architecture_manifest"] = architecture_hash

    all_events = load_jsonl(EVENT_PATH)
    live_events = [
        event for event in all_events
        if event.get("policy_mode") == "LIVE"
        and event.get("real_live_observation") is True
        and event.get("synthetic_initialization") is False
    ]
    historical_events = [event for event in all_events if event not in live_events]
    passive_records = load_jsonl(PASSIVE_PATH)
    passive_by_id: dict[str, dict[str, Any]] = {}
    for record in passive_records:
        identity = nested(record, "identity", "action_event_id")
        if not isinstance(identity, str) or not identity:
            continue
        if identity in passive_by_id:
            raise ValueError(f"Duplicate passive action_event_id: {identity}")
        passive_by_id[identity] = record
    missing_passive = [
        event["action_event_id"] for event in live_events
        if event["action_event_id"] not in passive_by_id
    ]
    if missing_passive:
        raise ValueError(f"LIVE events lack passive text/BKT enrichment: {missing_passive}")
    rows = [build_row(event, passive_by_id[event["action_event_id"]]) for event in live_events]

    validation = validate_inputs(state, lineage, architecture, live_events, rows)
    status_counts = Counter(str(row["resolution_status"]) for row in rows)
    move_counts = Counter(str(row["final_move"]) for row in rows)
    log_update_counts = {
        arm: sum(
            row["posterior_updated"] is True and row["selected_arm"] == arm
            for row in rows
        )
        for arm in ARMS
    }
    real_live_updates = sum(log_update_counts.values())
    initial_total = int(lineage["initial_total_update_count"])
    total_consistent = int(state["update_count"]) == initial_total + real_live_updates
    per_arm_consistent = all(
        int(state["arm_update_counts"][arm])
        == int(lineage["synthetic_arm_update_counts"][arm]) + log_update_counts[arm]
        for arm in ARMS
    )
    applied_ids = set(str(value) for value in state["applied_update_ids"])
    updated_event_ids = {
        str(row["action_event_id"]) for row in rows if row["posterior_updated"] is True
    }
    identities_consistent = updated_event_ids.issubset(applied_ids)
    accounting_status = (
        "CONSISTENT"
        if total_consistent and per_arm_consistent and identities_consistent
        else "STATE / LOG COUNT MISMATCH"
    )
    counts = {
        "total_live_actions": len(rows),
        "observed_update": status_counts["observed_update"],
        "resolved_no_update": status_counts["resolved_no_update"],
        "censored_no_response": status_counts["censored_no_response"],
        "processing_failed": status_counts["processing_failed"],
        "move_counts": {arm: move_counts[arm] for arm in ARMS},
        "real_live_updates": real_live_updates,
    }

    wb = Workbook()
    build_summary_sheet(wb, state, lineage, counts, source_hashes_before, accounting_status)
    write_grouped_table(
        wb.create_sheet("Decision_Only"),
        "Decision Only — Current LIVE Actions",
        "Rows show PRE-ACTION S+K+L inputs, the four sampled Turn-LinTS scores, the selected pedagogical move, and the resulting Tutor/learner/BKT feedback.",
        DECISION_COLUMNS,
        rows,
    )
    write_grouped_table(
        wb.create_sheet("LIVE_Turn_IO"),
        "LIVE Turn Input/Output — Forensic View",
        "Detailed current-LIVE action records joined to passive Tutor/learner/BKT evidence by action_event_id.",
        [*DECISION_COLUMNS, *FORENSIC_EXTRA_COLUMNS],
        rows,
    )
    build_field_guide(wb)
    build_architecture_flow(wb)
    build_posterior_sheet(wb, state, lineage, source_hashes_before["policy_state"], log_update_counts)
    build_update_audit(wb, rows, state, lineage, log_update_counts)
    build_historical_sheet(wb, historical_events)

    SCRIPT_DIR.mkdir(parents=True, exist_ok=True)
    temporary_workbook = WORKBOOK_PATH.with_name(f".{WORKBOOK_PATH.name}.tmp.xlsx")
    try:
        wb.save(temporary_workbook)
        check = load_workbook(temporary_workbook, read_only=True, data_only=False)
        required_sheets = {
            "Summary", "Decision_Only", "LIVE_Turn_IO", "Field_Guide",
            "Architecture_Flow", "Posterior_State", "LIVE_Update_Audit",
        }
        if not required_sheets.issubset(check.sheetnames):
            raise ValueError("Saved workbook is missing required sheets.")
        check.close()
        source_hashes_after = {key: sha256(path) for key, path in source_paths.items()}
        source_hashes_after["architecture_manifest"] = sha256(architecture_path)
        if source_hashes_after != source_hashes_before:
            raise RuntimeError(
                "A runtime source changed during generation; rerun for an atomic view."
            )
        os.replace(temporary_workbook, WORKBOOK_PATH)
    finally:
        temporary_workbook.unlink(missing_ok=True)

    validation["workbook_reopens_successfully"] = True
    validation["source_hashes_unchanged_during_generation"] = True
    validation["primary_live_rows"] = len(rows)
    summary = {
        "schema_version": "adaptmath_move_selector_architecture_io_view_v1",
        "source_runtime_lineage": relative(RUNTIME_ROOT),
        "source_event_log_path": relative(EVENT_PATH),
        "passive_enrichment_path": relative(PASSIVE_PATH),
        "policy_state_path": relative(STATE_PATH),
        "policy_state_sha256": source_hashes_before["policy_state"],
        "lineage_manifest_path": relative(LINEAGE_PATH),
        "architecture_manifest_path": relative(architecture_path),
        "LIVE_rows_exported": len(rows),
        "observed_updates": counts["observed_update"],
        "resolved_no_updates": counts["resolved_no_update"],
        "censored_rows": counts["censored_no_response"],
        "processing_failed_rows": counts["processing_failed"],
        "move_counts": counts["move_counts"],
        "current_state_update_count": int(state["update_count"]),
        "synthetic_initialization_count": int(lineage["synthetic_observation_count"]),
        "synthetic_effective_sample_size": float(lineage["synthetic_effective_sample_size"]),
        "derived_real_LIVE_update_count": real_live_updates,
        "real_LIVE_updates_by_arm": log_update_counts,
        "state_log_accounting": accounting_status,
        "state_log_discrepancy": None if accounting_status == "CONSISTENT" else {
            "total_consistent": total_consistent,
            "per_arm_consistent": per_arm_consistent,
            "updated_event_ids_present_in_state": identities_consistent,
        },
        "workbook_path": relative(WORKBOOK_PATH),
        "regeneration_script_path": relative(Path(__file__)),
        "sheets": wb.sheetnames,
        "validation": validation,
        "source_hashes": source_hashes_before,
    }
    temporary_summary = SUMMARY_PATH.with_name(f".{SUMMARY_PATH.name}.{os.getpid()}.tmp")
    temporary_summary.write_text(
        json.dumps(summary, indent=2, ensure_ascii=False, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    os.replace(temporary_summary, SUMMARY_PATH)
    print(json.dumps(summary, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
